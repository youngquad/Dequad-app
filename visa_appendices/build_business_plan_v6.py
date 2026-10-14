"""DEQUAD UKES Business Plan v6.0 — financial model + document builder.

Every number in the business plan's financial sections (11-18) is computed
here from the driver assumptions below, so the revenue forecast, P&L, cash
flow and balance sheet always reconcile with each other.

Run with:  python visa_appendices/build_business_plan_v6.py
Outputs (next to this script):
  - DEQUAD_UKES_Business_Plan_v6.md    (template tokens filled in)
  - DEQUAD_UKES_Business_Plan_v6.html / .pdf
  - DEQUAD_Financial_Model_v6.xlsx     (36-month workings, both scenarios)
"""
from __future__ import annotations

import math
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

HERE = Path(__file__).resolve().parent
TEMPLATE_MD = HERE / "business_plan_v6_template.md"
OUT_MD = HERE / "DEQUAD_UKES_Business_Plan_v6.md"
OUT_XLSX = HERE / "DEQUAD_Financial_Model_v6.xlsx"

MONTHS = 36
# M1 = September 2026. Financial year ends 31 August (aligned to the UK
# academic year, which is when universities buy and students sign up).
MONTH_NAMES = ["Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug"]


def label(m: int) -> str:
    year = 2026 + (m + 7) // 12  # m=1 -> Sep 2026, m=5 -> Jan 2027
    return f"{MONTH_NAMES[(m - 1) % 12]}-{str(year)[2:]}"


def yr(m: int) -> int:
    return (m - 1) // 12 + 1


# ---------------------------------------------------------------------------
# Assumptions
# ---------------------------------------------------------------------------
PREMIUM_PRICE = 4.99           # £/month, live in Stripe + RevenueCat
STARTER_PRICE = 49.99          # £/month, DEQUAD University Dashboard (self-serve, Stripe)
LICENCE_PRICE_PER_STUDENT = 2.00  # £/enrolled student/year, partner-university licence
SIGNUP_SEASONALITY = [0.16, 0.14, 0.10, 0.04, 0.12, 0.09, 0.07, 0.05, 0.04, 0.03, 0.03, 0.13]
ACTIVE_MONTHLY_ATTRITION = 0.06
OPENING_REGISTERED = 20        # closed-beta verified accounts at 1 Sep 2026
PREMIUM_LIVE_FROM = 3          # paywall + live billing from Nov 2026
PREMIUM_FEE_RATE = 0.10        # blended Stripe (~5.5% at £4.99) and App Store/Play (15%)
STARTER_FEE_RATE = 0.019       # Stripe 1.5% + 20p on £49.99
CORP_TAX = 0.19                # small profits rate (profits <= £50k)
NI_RATE, NI_THRESHOLD, EMPLOYMENT_ALLOWANCE = 0.15, 5_000, 10_500
EMPLOYER_PENSION = 0.03        # paid on all salaries, regardless of auto-enrolment thresholds
NATWEST_IN_KIND_UNTIL = 7      # office, legal and accountancy in kind until Mar 2027
VAT_THRESHOLD = 90_000


@dataclass
class Licence:
    name: str
    students: int
    start: int  # month the 12-month term starts
    renews: bool = True

    @property
    def annual_value(self) -> float:
        return self.students * LICENCE_PRICE_PER_STUDENT

    def active(self, m: int) -> bool:
        if m < self.start:
            return False
        return self.renews or m < self.start + 12

    def invoiced(self, m: int) -> float:
        # Invoiced on signature / renewal one month before each term starts,
        # paid on 30-day terms, so cash lands in the first month of the term.
        n = m + 1
        if n < self.start:
            return 0.0
        if (n - self.start) % 12 == 0 and (self.renews or n == self.start):
            return self.annual_value
        return 0.0


@dataclass
class Hire:
    role: str
    start: int
    annual_salary: float
    fte: float
    director: bool = False


@dataclass
class Scenario:
    name: str
    annual_new_signups: tuple[int, int, int]
    premium_conversion: tuple[float, float, float]
    licensed_share_of_active: tuple[float, float, float]  # share of active base whose uni holds a licence
    starter_adds: dict[int, int]
    starter_churn: dict[int, int]
    licences: list[Licence]
    icb_contracts: list[tuple[int, float]]  # (start month, annual value)
    staff: list[Hire]
    marketing: dict[str, tuple[int, int, int]]
    capex: dict[int, list[tuple[str, float]]]
    pilot_cost: dict[int, float] = field(default_factory=lambda: {5: 300.0})
    lean: bool = False  # downside test: overheads held at Year 1 run-rate, remote working


BASE_MARKETING = {
    "Student ambassador programme (incentives, referral rewards)": (300, 900, 1_500),
    "Freshers' / Welcome Week stands & printed materials": (200, 700, 1_000),
    "Instagram & TikTok (organic content + boosted posts)": (150, 700, 1_200),
    "Google Search ads (brand + 'student friends' terms)": (0, 300, 600),
    "Content, SEO & safeguarding white paper": (50, 300, 500),
    "B2B events & PR (AMOSSHE, UUK, Wonkhe, NatWest network)": (100, 600, 1_200),
}

BASE = Scenario(
    name="Base case",
    annual_new_signups=(300, 1_500, 4_000),
    premium_conversion=(0.025, 0.035, 0.045),
    licensed_share_of_active=(0.0, 0.50, 0.60),
    starter_adds={8: 1, 11: 1, 13: 1, 15: 1, 18: 1, 21: 1, 25: 1, 27: 1, 29: 1, 31: 1, 33: 1, 35: 1},
    starter_churn={30: 1},
    licences=[
        Licence("Partner licence 1 (anchor university, c.10,000 students)", 10_000, 14),
        Licence("Partner licence 2 (c.10,000 students)", 10_000, 25),
        Licence("Partner licence 3 (c.6,000 students)", 6_000, 31),
    ],
    icb_contracts=[],
    staff=[
        Hire("Founder & CEO — Yusuf Quadri", 25, 10_000, 1.0, director=True),
        Hire("Co-founder & CTO — Yusuff Adeagbo", 25, 10_000, 1.0, director=True),
        Hire("Safeguarding & Trust Lead (part-time, c.10 hrs/wk)", 25, 8_000, 0.25),
    ],
    marketing=BASE_MARKETING,
    capex={
        16: [("QA test handsets (1 x iOS, 1 x Android, refurbished)", 600.0),
             ("Development laptop (CTO replacement)", 900.0)],
        26: [("Laptop — Safeguarding & Trust Lead", 850.0), ("Campus event kit (banner, display)", 350.0)],
    },
)

GROWTH = Scenario(
    name="Growth case",
    annual_new_signups=(300, 2_500, 9_000),
    premium_conversion=(0.025, 0.035, 0.045),
    licensed_share_of_active=(0.0, 0.55, 0.70),
    starter_adds={8: 1, 11: 1, 13: 1, 14: 1, 15: 1, 17: 1, 18: 1, 21: 1, 25: 1, 26: 1, 27: 1, 28: 1,
                  29: 1, 31: 1, 33: 1, 35: 1},
    starter_churn={30: 1},
    licences=[
        Licence("Partner licence 1", 10_000, 14),
        Licence("Partner licence 2", 10_000, 17),
        Licence("Partner licence 3", 12_000, 25),
        Licence("Partner licence 4", 10_000, 25),
        Licence("Partner licence 5", 8_000, 26),
        Licence("Partner licence 6", 15_000, 28),
        Licence("Partner licence 7", 6_000, 31),
    ],
    icb_contracts=[(31, 20_000.0)],
    staff=[
        Hire("Founder & CEO — Yusuf Quadri", 25, 20_000, 1.0, director=True),
        Hire("Co-founder & CTO — Yusuff Adeagbo", 25, 20_000, 1.0, director=True),
        Hire("Safeguarding & Trust Lead (part-time)", 25, 8_000, 0.25),
        Hire("Sales Lead (full-time)", 25, 25_000, 1.0),
        Hire("Support Agent (full-time)", 26, 25_000, 1.0),
        Hire("Student Ambassador (full-time)", 26, 25_000, 1.0),
    ],
    marketing={k: (v[0], round(v[1] * 1.5, -1), round(v[2] * 2.0, -1)) for k, v in BASE_MARKETING.items()},
    capex={
        16: [("QA test handsets", 600.0), ("Development laptop", 900.0)],
        25: [("Laptops x 3 (new hires)", 2_550.0), ("Campus event kit", 350.0)],
        26: [("Laptop (Support Agent)", 850.0)],
    },
)


# ---------------------------------------------------------------------------
# Engine
# ---------------------------------------------------------------------------
def run(s: Scenario) -> dict:
    r: dict[str, list[float]] = {k: [0.0] * (MONTHS + 1) for k in [
        "new_signups", "registered", "active", "premium_subs", "starter_subs", "licences_active",
        "licence_students", "rev_premium", "rev_starter", "rev_licence", "rev_icb", "revenue",
        "cos_hosting", "cos_ai", "cos_messaging", "cos_payment", "cos_onboarding", "cos_success",
        "cos_total", "gross_profit",
        "salaries", "employer_ni", "pension", "other_staff", "staff_total",
        "marketing", "software", "insurance", "accountancy", "legal", "office", "compliance",
        "travel", "bank_comms", "overheads_ex_staff", "ebitda", "depreciation", "operating_profit",
        "capex", "receipts", "licence_invoiced", "payments", "ct_paid", "net_cash", "closing_cash",
        "headcount", "fte",
    ]}
    registered = OPENING_REGISTERED
    active = OPENING_REGISTERED * 0.5
    starter = 0
    cash = 6_000.0  # founder share capital, paid into the company bank account at incorporation
    deferred = 0.0
    debtors = 0.0
    licence_invoices_pending: dict[int, float] = {}
    assets: list[tuple[int, float]] = []
    for m in range(1, MONTHS + 1):
        y = yr(m)
        new = s.annual_new_signups[y - 1] * SIGNUP_SEASONALITY[(m - 1) % 12]
        registered += new
        active = active * (1 - ACTIVE_MONTHLY_ATTRITION) + new
        uncovered = active * (1 - s.licensed_share_of_active[y - 1])
        premium = round(uncovered * s.premium_conversion[y - 1]) if m >= PREMIUM_LIVE_FROM else 0
        starter += s.starter_adds.get(m, 0) - s.starter_churn.get(m, 0)
        lic_active = [l for l in s.licences if l.active(m)]
        r["new_signups"][m] = new
        r["registered"][m] = registered
        r["active"][m] = active
        r["premium_subs"][m] = premium
        r["starter_subs"][m] = starter
        r["licences_active"][m] = len(lic_active)
        r["licence_students"][m] = sum(l.students for l in lic_active)

        rev_p = premium * PREMIUM_PRICE
        rev_s = starter * STARTER_PRICE
        rev_l = sum(l.annual_value / 12 for l in lic_active)
        rev_i = sum(v / 12 for start, v in s.icb_contracts if m >= start)
        revenue = rev_p + rev_s + rev_l + rev_i
        r["rev_premium"][m], r["rev_starter"][m], r["rev_licence"][m], r["rev_icb"][m] = rev_p, rev_s, rev_l, rev_i
        r["revenue"][m] = revenue

        new_licences = sum(1 for l in s.licences if l.start == m) + sum(1 for st, _ in s.icb_contracts if st == m)
        cos = {
            "cos_hosting": 35 + 0.02 * registered,
            "cos_ai": 20 + 0.01 * active + 40 * (len(lic_active) + sum(1 for st, _ in s.icb_contracts if m >= st)),
            "cos_messaging": 10 + 0.005 * registered,
            "cos_payment": rev_p * PREMIUM_FEE_RATE + rev_s * STARTER_FEE_RATE,
            "cos_onboarding": 600 * new_licences + s.pilot_cost.get(m, 0),
            "cos_success": 20 * len(lic_active),
        }
        for k, v in cos.items():
            r[k][m] = v
        r["cos_total"][m] = sum(cos.values())
        r["gross_profit"][m] = revenue - r["cos_total"][m]

        # Staff: NI computed on annualised salary above the secondary threshold, then
        # offset by the Employment Allowance (available because at least one
        # non-director employee is on payroll from the same month as the directors).
        sal = ni = 0.0
        hc = 2  # the two founders are always on the team, paid or not
        fte = 2.0
        for h in s.staff:
            if m >= h.start:
                sal += h.annual_salary / 12
                ni += max(0.0, h.annual_salary - NI_THRESHOLD) * NI_RATE / 12
                if not h.director:
                    hc += 1
                    fte += h.fte
        has_non_director = any(not h.director and m >= h.start for h in s.staff)
        r["salaries"][m] = sal
        r["employer_ni"][m] = ni  # gross; EA offset applied annually below
        r["pension"][m] = sal * EMPLOYER_PENSION
        r["other_staff"][m] = sum(60 + 150 for h in s.staff if h.start == m and not h.director)  # DBS + induction/training
        r["headcount"][m] = hc
        r["fte"][m] = fte
        r["_has_non_director"] = has_non_director

        idx = (m - 1) % 12
        mk_total = sum(v[y - 1] for v in s.marketing.values())
        r["marketing"][m] = mk_total * SIGNUP_SEASONALITY[idx]
        yc = 1 if s.lean else y  # lean: every overhead held at its Year 1 run-rate
        r["software"][m] = {1: 45, 2: 90, 3: 130}[yc] + (79 + 20 if m == 1 else 79 if m in (13, 25) else 0)
        r["insurance"][m] = {1: 600, 2: 1_100, 3: 1_750 + (600 if s is GROWTH else 0)}[yc] / 12  # monthly instalments
        r["accountancy"][m] = {16: 600, 28: 750}.get(m, 0) + (0 if yc < 3 else 40 * (2 if s is GROWTH else 1))
        r["legal"][m] = 0 if s.lean else {13: 400, 24: 500, 30: 400}.get(m, 0)
        r["office"][m] = 0 if m <= NATWEST_IN_KIND_UNTIL or s.lean else {1: 25, 2: 50, 3: 100}[y] * (3 if s is GROWTH and y == 3 else 1)
        r["compliance"][m] = (52 if m in (1, 13, 25) else 0) + (384 if m in (4, 16, 28) else 0) + \
            (50 if m in (12, 24, 36) else 0) + (270 if m == 2 else 0)
        r["travel"][m] = {1: 30, 2: 100, 3: 150}[yc] * (2 if s is GROWTH and y == 3 else 1)
        r["bank_comms"][m] = 10 + (0 if yc == 1 else 8)

        overheads = sum(r[k][m] for k in ["marketing", "software", "insurance", "accountancy", "legal",
                                           "office", "compliance", "travel", "bank_comms"])
        r["overheads_ex_staff"][m] = overheads

        for desc, cost in s.capex.get(m, []):
            assets.append((m, cost))
        r["capex"][m] = sum(c for _, c in s.capex.get(m, []))
        r["depreciation"][m] = sum(c / 36 for start, c in assets if start <= m < start + 36)

        # Cash: licences and ICB contracts are invoiced annually in advance, paid
        # on 30-day terms (cash the month after invoice).
        invoiced = sum(l.invoiced(m) for l in s.licences) + \
            sum(v for st, v in s.icb_contracts if m + 1 >= st and (m + 1 - st) % 12 == 0)
        r["licence_invoiced"][m] = invoiced
        licence_invoices_pending[m + 1] = licence_invoices_pending.get(m + 1, 0) + invoiced
        licence_cash = licence_invoices_pending.pop(m, 0.0)
        r["receipts"][m] = rev_p + rev_s + licence_cash
        debtors += invoiced - licence_cash
        deferred += invoiced - rev_l - rev_i

    # Employment Allowance: offset employer NI per tax year (approximated by FY).
    for y in (1, 2, 3):
        ms = range((y - 1) * 12 + 1, y * 12 + 1)
        gross = sum(r["employer_ni"][m] for m in ms)
        if gross and any(not h.director and (y - 1) * 12 < h.start <= y * 12 or (not h.director and h.start <= (y - 1) * 12) for h in s.staff):
            allowance = min(gross, EMPLOYMENT_ALLOWANCE)
            for m in ms:
                r["employer_ni"][m] -= r["employer_ni"][m] / gross * allowance
        r.setdefault("_ni_gross", {})[y] = gross

    for m in range(1, MONTHS + 1):
        r["staff_total"][m] = r["salaries"][m] + r["employer_ni"][m] + r["pension"][m] + r["other_staff"][m]
        r["ebitda"][m] = r["gross_profit"][m] - r["staff_total"][m] - r["overheads_ex_staff"][m]
        r["operating_profit"][m] = r["ebitda"][m] - r["depreciation"][m]

    # Corporation tax: taxable = operating profit + depreciation - capex (100% AIA),
    # with trading losses carried forward. Paid 9 months + 1 day after year end
    # (FY ends 31 Aug -> paid 1 June -> month 10 of the following year).
    annual = {}
    losses = 0.0
    tax_due = {}
    for y in (1, 2, 3):
        ms = range((y - 1) * 12 + 1, y * 12 + 1)
        op = sum(r["operating_profit"][m] for m in ms)
        taxable = op + sum(r["depreciation"][m] for m in ms) - sum(r["capex"][m] for m in ms)
        if taxable < 0:
            losses += -taxable
            tax = 0.0
        else:
            used = min(losses, taxable)
            losses -= used
            tax = round((taxable - used) * CORP_TAX, 0)
        tax_due[y] = tax
        annual[y] = {"taxable_before_losses": taxable, "tax": tax, "losses_cf": losses}
    ct_payment_month = {1: 22, 2: 34}
    cash = 6_000.0
    for m in range(1, MONTHS + 1):
        ct = sum(tax_due[y] for y, pm in ct_payment_month.items() if pm == m)
        r["ct_paid"][m] = ct
        pay = r["cos_total"][m] + r["staff_total"][m] + r["overheads_ex_staff"][m] + r["capex"][m] + ct
        r["payments"][m] = pay
        r["net_cash"][m] = r["receipts"][m] - pay
        cash += r["net_cash"][m]
        r["closing_cash"][m] = cash

    # Year-end balance sheets
    bs = {}
    retained = 0.0
    for y in (1, 2, 3):
        end = y * 12
        ms = range((y - 1) * 12 + 1, end + 1)
        profit_after_tax = sum(r["operating_profit"][m] for m in ms) - tax_due[y]
        retained += profit_after_tax
        cost = sum(c for mo, lst in s.capex.items() if mo <= end for _, c in lst)
        acc_dep = sum(r["depreciation"][m] for m in range(1, end + 1))
        invoiced_to_date = sum(r["licence_invoiced"][m] for m in range(1, end + 1))
        licence_cash_to_date = invoiced_to_date - r["licence_invoiced"][end]  # 30-day terms
        debt = invoiced_to_date - licence_cash_to_date
        recognised = sum(r["rev_licence"][m] + r["rev_icb"][m] for m in range(1, end + 1))
        deferred_income = invoiced_to_date - recognised
        ct_creditor = sum(tax_due[yy] for yy in (1, 2, 3) if yy <= y and ct_payment_month.get(yy, 99) > end)
        assets_total = (cost - acc_dep) + r["closing_cash"][end] + debt
        liabilities = deferred_income + ct_creditor
        bs[y] = {
            "fixed_cost": cost, "acc_dep": acc_dep, "nbv": cost - acc_dep, "cash": r["closing_cash"][end],
            "debtors": debt, "stock": 0.0, "total_assets": assets_total, "deferred_income": deferred_income,
            "ct_creditor": ct_creditor, "total_liabilities": liabilities, "net_assets": assets_total - liabilities,
            "share_capital": 6_000.0, "retained": retained, "equity": 6_000.0 + retained,
        }
        assert abs(bs[y]["net_assets"] - bs[y]["equity"]) < 0.5, (s.name, y, bs[y])
    r["_annual_tax"] = annual
    r["_tax_due"] = tax_due
    r["_bs"] = bs
    return r


def ysum(r, key, y):
    return sum(r[key][m] for m in range((y - 1) * 12 + 1, y * 12 + 1))


def yavg(r, key, y):
    return ysum(r, key, y) / 12


def yend(r, key, y):
    return r[key][y * 12]


# ---------------------------------------------------------------------------
# Formatting
# ---------------------------------------------------------------------------
def gbp(v: float, bold=False) -> str:
    v = round(v)
    s = f"£{abs(v):,}" if v >= 0 else f"(£{abs(v):,})"
    if v == 0:
        s = "–"
    return f"**{s}**" if bold else s


def num(v: float, dp=0) -> str:
    return f"{v:,.{dp}f}"


def pct(v: float) -> str:
    return f"{v * 100:.1f}%"


def table(header: list[str], rows: list[list[str]], align_right_from=1) -> str:
    out = ["| " + " | ".join(header) + " |",
           "|" + "|".join(["---"] * align_right_from + ["---:"] * (len(header) - align_right_from)) + "|"]
    out += ["| " + " | ".join(row) + " |" for row in rows]
    return "\n".join(out)


def annual_row(r, label_, key, bold=False, neg=False):
    vals = [ysum(r, key, y) * (-1 if neg else 1) for y in (1, 2, 3)]
    lbl = f"**{label_}**" if bold else label_
    return [lbl] + [gbp(v, bold) for v in vals] + [gbp(sum(vals), bold)]


def build_tokens(b: dict, g: dict) -> dict[str, str]:
    t: dict[str, str] = {}
    Y = (1, 2, 3)
    H3 = ["", "Year 1 (FY27)", "Year 2 (FY28)", "Year 3 (FY29)"]
    H3T = H3 + ["3-yr total"]

    # Headline values
    for y in Y:
        t[f"REV_Y{y}"] = gbp(ysum(b, "revenue", y))
        t[f"GP_Y{y}"] = gbp(ysum(b, "gross_profit", y))
        t[f"GM_Y{y}"] = pct(ysum(b, "gross_profit", y) / ysum(b, "revenue", y)) if ysum(b, "revenue", y) > 1000 else "n/a"
        t[f"EBITDA_Y{y}"] = gbp(ysum(b, "ebitda", y))
        t[f"OP_Y{y}"] = gbp(ysum(b, "operating_profit", y))
        t[f"PAT_Y{y}"] = gbp(ysum(b, "operating_profit", y) - b["_tax_due"][y])
        t[f"CASH_Y{y}"] = gbp(yend(b, "closing_cash", y))
        t[f"REG_Y{y}"] = num(yend(b, "registered", y))
        t[f"ACTIVE_Y{y}"] = num(yend(b, "active", y))
        t[f"PREM_END_Y{y}"] = num(yend(b, "premium_subs", y))
        t[f"STARTER_END_Y{y}"] = num(yend(b, "starter_subs", y))
        t[f"LIC_END_Y{y}"] = num(yend(b, "licences_active", y))
        t[f"LICSTUD_END_Y{y}"] = num(yend(b, "licence_students", y))
        t[f"MKT_Y{y}"] = gbp(ysum(b, "marketing", y))
        t[f"STAFF_Y{y}"] = gbp(ysum(b, "staff_total", y))
        t[f"G_REV_Y{y}"] = gbp(ysum(g, "revenue", y))
        t[f"G_CASH_Y{y}"] = gbp(yend(g, "closing_cash", y))
        t[f"G_OP_Y{y}"] = gbp(ysum(g, "operating_profit", y))
    t["MIN_CASH"] = gbp(min(b["closing_cash"][1:]))
    t["MIN_CASH_MONTH"] = label(min(range(1, MONTHS + 1), key=lambda m: b["closing_cash"][m]))
    t["G_MIN_CASH"] = gbp(min(g["closing_cash"][1:]))
    t["G_MIN_CASH_MONTH"] = label(min(range(1, MONTHS + 1), key=lambda m: g["closing_cash"][m]))
    t["Y1_BURN"] = gbp(-ysum(b, "net_cash", 1))
    t["REV_3YR"] = gbp(sum(ysum(b, "revenue", y) for y in Y))
    t["G_PEAK_HC"] = num(max(g["headcount"][1:]))
    t["G_FTE_END"] = num(g["fte"][36], 2)
    t["FTE_END"] = num(b["fte"][36], 2)
    t["HC_END"] = num(b["headcount"][36])

    # 11.1 Customer acquisition
    rows = []
    for lbl, key, fn in [
        ("New verified student sign-ups in year", "new_signups", "sum"),
        ("Cumulative verified students (year end)", "registered", "end"),
        ("Monthly active students (year end)", "active", "end"),
        ("Premium subscribers (year end)", "premium_subs", "end"),
        ("Premium subscribers (average in year)", "premium_subs", "avg"),
        ("Insights Starter subscriptions (year end)", "starter_subs", "end"),
        ("Insights Starter subscriptions (average)", "starter_subs", "avg"),
        ("Partner-university licences (year end)", "licences_active", "end"),
        ("Students covered by partner licences (year end)", "licence_students", "end"),
    ]:
        f = {"sum": ysum, "end": yend, "avg": yavg}[fn]
        rows.append([lbl] + [num(f(b, key, y), 1 if fn == "avg" else 0) for y in Y])
    t["TABLE_ACQUISITION"] = table(H3, rows)

    # Quarterly acquisition view
    qrows = []
    for q in range(12):
        ms = range(q * 3 + 1, q * 3 + 4)
        qrows.append([f"Y{q // 4 + 1} Q{q % 4 + 1} ({label(ms[0])} – {label(ms[-1])})",
                      num(sum(b["new_signups"][m] for m in ms)),
                      num(b["registered"][ms[-1]]),
                      num(b["premium_subs"][ms[-1]]),
                      num(b["starter_subs"][ms[-1]]),
                      num(b["licences_active"][ms[-1]]),
                      gbp(sum(b["revenue"][m] for m in ms))])
    t["TABLE_ACQ_QUARTERLY"] = table(["Quarter", "New sign-ups", "Verified students", "Premium subs",
                                      "Starter subs", "Licences", "Revenue"], qrows)

    # 11.2 Revenue
    rows = [annual_row(b, "Partner-university licences (£2/student/yr)", "rev_licence"),
            annual_row(b, "Insights Starter (£49.99/month)", "rev_starter"),
            annual_row(b, "DEQUAD Premium (£4.99/month)", "rev_premium"),
            annual_row(b, "NHS ICB contracts", "rev_icb"),
            annual_row(b, "Total revenue", "revenue", bold=True)]
    t["TABLE_REVENUE"] = table(H3T, rows)

    rows = [annual_row(b, "Cloud hosting & database", "cos_hosting"),
            annual_row(b, "AI inference (safeguarding triage, cohort analysis)", "cos_ai"),
            annual_row(b, "Email, push & SMS", "cos_messaging"),
            annual_row(b, "Payment processing & app-store fees", "cos_payment"),
            annual_row(b, "University onboarding & pilot delivery", "cos_onboarding"),
            annual_row(b, "Customer-success tooling", "cos_success"),
            annual_row(b, "Total cost of sales", "cos_total", bold=True),
            annual_row(b, "Gross profit", "gross_profit", bold=True)]
    rows.append(["**Gross margin**"] + [f"**{t[f'GM_Y{y}']}**" for y in Y] +
                [f"**{pct(sum(ysum(b, 'gross_profit', y) for y in Y) / sum(ysum(b, 'revenue', y) for y in Y))}**"])
    t["TABLE_COS"] = table(H3T, rows)

    # 12 Cash flow — annual
    def cf_rows(r):
        rows = [["Opening cash"] + [gbp(6_000 if y == 1 else yend(r, "closing_cash", y - 1)) for y in Y] + [gbp(6_000)]]
        rows.append(annual_row(r, "Receipts — Premium & Starter", "rev_premium"))
        rows[-1] = ["Receipts — Premium & Starter subscriptions"] + [
            gbp(ysum(r, "rev_premium", y) + ysum(r, "rev_starter", y)) for y in Y] + [
            gbp(sum(ysum(r, "rev_premium", y) + ysum(r, "rev_starter", y) for y in Y))]
        lic = [ysum(r, "receipts", y) - ysum(r, "rev_premium", y) - ysum(r, "rev_starter", y) for y in Y]
        rows.append(["Receipts — licences & contracts (annual in advance)"] + [gbp(v) for v in lic] + [gbp(sum(lic))])
        rows.append(annual_row(r, "Total receipts", "receipts", bold=True))
        rows.append(annual_row(r, "Cost of sales", "cos_total", neg=True))
        rows.append(annual_row(r, "Staff costs (salaries, NI, pension)", "staff_total", neg=True))
        rows.append(annual_row(r, "Marketing", "marketing", neg=True))
        other = [ysum(r, "overheads_ex_staff", y) - ysum(r, "marketing", y) for y in Y]
        rows.append(["Other overheads (software, insurance, professional, office, compliance, travel)"] +
                    [gbp(-v) for v in other] + [gbp(-sum(other))])
        rows.append(annual_row(r, "Capital expenditure", "capex", neg=True))
        rows.append(annual_row(r, "Corporation tax paid", "ct_paid", neg=True))
        rows.append(annual_row(r, "Total payments", "payments", bold=True, neg=True))
        rows.append(annual_row(r, "Net cash flow", "net_cash", bold=True))
        rows.append(["**Closing cash**"] + [gbp(yend(r, "closing_cash", y), True) for y in Y] +
                    [gbp(yend(r, "closing_cash", 3), True)])
        return rows
    t["TABLE_CASH_ANNUAL"] = table(H3T, cf_rows(b))
    t["TABLE_G_CASH_ANNUAL"] = table(H3T, cf_rows(g))

    # Monthly Y1
    def monthly(r, ms):
        hdr = ["£"] + [label(m) for m in ms] + ["Total"]
        def line(name, key, neg=False, bold=False):
            vals = [r[key][m] * (-1 if neg else 1) for m in ms]
            n = f"**{name}**" if bold else name
            return [n] + [gbp(v, bold) for v in vals] + [gbp(sum(vals), bold)]
        rows = [["Opening cash"] + [gbp(r["closing_cash"][m - 1] if m > 1 else 6_000) for m in ms] + [""]]
        rows += [line("Receipts", "receipts", bold=True),
                 line("Cost of sales", "cos_total", neg=True),
                 line("Staff", "staff_total", neg=True),
                 line("Marketing", "marketing", neg=True)]
        oth = [-(r["overheads_ex_staff"][m] - r["marketing"][m]) for m in ms]
        rows.append(["Other overheads"] + [gbp(v) for v in oth] + [gbp(sum(oth))])
        rows += [line("Capex", "capex", neg=True), line("Corp. tax", "ct_paid", neg=True),
                 line("Net cash flow", "net_cash", bold=True)]
        rows.append(["**Closing cash**"] + [gbp(r["closing_cash"][m], True) for m in ms] + [""])
        return table(hdr, rows)
    t["TABLE_CASH_MONTHLY_Y1"] = monthly(b, range(1, 13))
    t["TABLE_CASH_MONTHLY_Y2"] = monthly(b, range(13, 25))
    t["TABLE_CASH_MONTHLY_Y3"] = monthly(b, range(25, 37))

    # 13 P&L
    def pl_rows(r):
        rows = [annual_row(r, "Revenue", "revenue", bold=True),
                annual_row(r, "Cost of sales", "cos_total", neg=True),
                annual_row(r, "Gross profit", "gross_profit", bold=True),
                annual_row(r, "Salaries", "salaries", neg=True),
                annual_row(r, "Employer NI (after Employment Allowance)", "employer_ni", neg=True),
                annual_row(r, "Employer pension (3%)", "pension", neg=True),
                annual_row(r, "Other staff costs (DBS, induction)", "other_staff", neg=True)]
        for k, lbl in [("marketing", "Advertising & marketing"), ("software", "Software & developer accounts"),
                       ("insurance", "Insurance (PI, cyber, public & employers' liability)"),
                       ("accountancy", "Accountancy & bookkeeping"), ("legal", "Legal"),
                       ("office", "Office / co-working"), ("compliance", "Compliance (ICO, Cyber Essentials, IP, filings)"),
                       ("travel", "Travel & campus visits"), ("bank_comms", "Bank, telephone & internet")]:
            rows.append(annual_row(r, lbl, k, neg=True))
        tot = [ysum(r, "staff_total", y) + ysum(r, "overheads_ex_staff", y) for y in Y]
        rows.append(["**Total overheads**"] + [gbp(-v, True) for v in tot] + [gbp(-sum(tot), True)])
        rows.append(annual_row(r, "EBITDA", "ebitda", bold=True))
        rows.append(annual_row(r, "Depreciation", "depreciation", neg=True))
        rows.append(annual_row(r, "Operating profit / (loss)", "operating_profit", bold=True))
        tax = [r["_tax_due"][y] for y in Y]
        rows.append(["Corporation tax (19%, after losses b/f and 100% AIA)"] + [gbp(-v) for v in tax] + [gbp(-sum(tax))])
        pat = [ysum(r, "operating_profit", y) - r["_tax_due"][y] for y in Y]
        rows.append(["**Profit / (loss) after tax**"] + [gbp(v, True) for v in pat] + [gbp(sum(pat), True)])
        rows.append(["Net margin"] + [pct(p / ysum(r, "revenue", y)) if ysum(r, "revenue", y) > 1000 else "n/a"
                                      for p, y in zip(pat, Y)] + [""])
        return rows
    t["TABLE_PL"] = table(H3T, pl_rows(b))
    t["TABLE_G_PL"] = table(H3T, pl_rows(g))

    # Tax computation
    rows = []
    for y in Y:
        a = b["_annual_tax"][y]
        rows.append([f"Year {y}", gbp(ysum(b, "operating_profit", y)), gbp(ysum(b, "depreciation", y)),
                     gbp(-ysum(b, "capex", y)), gbp(a["taxable_before_losses"]), gbp(a["losses_cf"]), gbp(a["tax"]),
                     {1: label(22), 2: label(34), 3: "Jun-30 (after forecast)"}[y]])
    t["TABLE_TAX"] = table(["", "Operating result", "Add back depreciation", "Less AIA", "Tax-adjusted result",
                            "Losses c/f", "CT @19%", "Paid"], rows)

    # 14 Balance sheet
    def bs_rows(r):
        bs = r["_bs"]
        out = []
        for lbl, k, bold in [("Fixed assets at cost", "fixed_cost", False), ("Less accumulated depreciation", "acc_dep", False),
                             ("**Net book value of fixed assets**", "nbv", True), ("Stock", "stock", False),
                             ("Trade debtors (licence invoices, 30-day terms)", "debtors", False),
                             ("Cash at bank", "cash", False), ("**Total assets**", "total_assets", True),
                             ("Deferred income (licence fees invoiced in advance)", "deferred_income", False),
                             ("Corporation tax payable", "ct_creditor", False),
                             ("**Total liabilities**", "total_liabilities", True), ("**Net assets**", "net_assets", True),
                             ("Share capital (founders, £3,000 each)", "share_capital", False),
                             ("Profit & loss reserve", "retained", False), ("**Total shareholders' funds**", "equity", True)]:
            sign = -1 if k in ("acc_dep", "deferred_income", "ct_creditor", "total_liabilities") else 1
            out.append([lbl] + [gbp(sign * bs[y][k], bold) for y in Y])
        return out
    t["TABLE_BS"] = table(["At 31 August", "2027 (Y1)", "2028 (Y2)", "2029 (Y3)"], bs_rows(b))

    # 15 Stock
    t["TABLE_STOCK"] = table(["", "Y1", "Y2", "Y3"], [["Opening stock", "–", "–", "–"], ["Purchases", "–", "–", "–"],
                                                     ["Closing stock", "–", "–", "–"]])

    # 16 Marketing
    rows = []
    for k, v in BASE_MARKETING.items():
        rows.append([k] + [gbp(x) for x in v] + [gbp(sum(v))])
    tot = [ysum(b, "marketing", y) for y in Y]
    rows.append(["**Total advertising & marketing**"] + [gbp(v, True) for v in tot] + [gbp(sum(tot), True)])
    rows.append(["Marketing as % of revenue"] + [pct(tot[i] / ysum(b, "revenue", y)) if ysum(b, "revenue", y) > 5000 else "n/a"
                                                 for i, y in enumerate(Y)] + [""])
    new_students = [ysum(b, "new_signups", y) for y in Y]
    rows.append(["Blended cost per verified student sign-up"] + [f"£{tot[i] / new_students[i]:.2f}" for i in range(3)] + [""])
    t["TABLE_MARKETING"] = table(["Channel"] + H3T[1:], rows)
    q = []
    for y in Y:
        ms = list(range((y - 1) * 12 + 1, y * 12 + 1))
        q.append([f"Year {y}"] + [gbp(sum(b["marketing"][m] for m in ms[i * 3:(i + 1) * 3])) for i in range(4)] +
                 [gbp(sum(b["marketing"][m] for m in ms))])
    t["TABLE_MARKETING_Q"] = table(["", "Q1 Sep–Nov", "Q2 Dec–Feb", "Q3 Mar–May", "Q4 Jun–Aug", "Total"], q)

    # 17 Fixed assets
    rows = []
    for m, items in sorted(BASE.capex.items()):
        for desc, cost in items:
            dep = [min(cost / 36 * 12, max(0.0, cost - cost / 36 * 12 * 0)) for _ in Y]
            dep_y = []
            for y in Y:
                ms = range((y - 1) * 12 + 1, y * 12 + 1)
                dep_y.append(sum(cost / 36 for mm in ms if m <= mm < m + 36))
            nbv = cost - sum(dep_y)
            rows.append([desc, label(m), gbp(cost)] + [gbp(d) for d in dep_y] + [gbp(nbv)])
    bs = b["_bs"]
    rows.append(["**Total**", "", gbp(sum(c for l in BASE.capex.values() for _, c in l), True)] +
                [gbp(ysum(b, "depreciation", y), True) for y in Y] + [gbp(bs[3]["nbv"], True)])
    t["TABLE_FIXED_ASSETS"] = table(["Asset", "Acquired", "Cost", "Dep'n Y1", "Dep'n Y2", "Dep'n Y3", "NBV at Y3 end"], rows, 2)
    rows = []
    for y in Y:
        opening = bs[y - 1]["nbv"] if y > 1 else 0
        rows.append([f"Year {y}", gbp(opening), gbp(ysum(b, "capex", y)), gbp(-ysum(b, "depreciation", y)), gbp(bs[y]["nbv"])])
    t["TABLE_FA_ROLL"] = table(["", "Opening NBV", "Additions", "Depreciation", "Closing NBV"], rows)

    # 18 Staff
    rows = []
    for h in BASE.staff:
        months_y3 = max(0, 36 - h.start + 1)
        sal = h.annual_salary * months_y3 / 12
        ni = max(0, h.annual_salary - NI_THRESHOLD) * NI_RATE * months_y3 / 12
        rows.append([h.role, label(h.start), f"{h.fte:.2f}", gbp(0), gbp(0), gbp(sal), gbp(ni), gbp(sal * EMPLOYER_PENSION)])
    t["TABLE_STAFF_ROLES"] = table(["Role", "Paid from", "FTE", "Y1 salary", "Y2 salary", "Y3 salary",
                                    "Y3 employer NI (gross)", "Y3 pension 3%"], rows, 2)
    rows = [annual_row(b, "Gross salaries", "salaries")]
    gross_ni = [b["_ni_gross"][y] for y in Y]
    rows.append(["Employer NI before allowance (15% above £5,000)"] + [gbp(v) for v in gross_ni] + [gbp(sum(gross_ni))])
    ea = [gross_ni[i] - ysum(b, "employer_ni", y) for i, y in enumerate(Y)]
    rows.append(["Less Employment Allowance"] + [gbp(-v) for v in ea] + [gbp(-sum(ea))])
    rows.append(annual_row(b, "Employer pension (3% of salary)", "pension"))
    rows.append(annual_row(b, "Other staff costs (enhanced DBS, induction & safeguarding training)", "other_staff"))
    rows.append(annual_row(b, "Total staff cost", "staff_total", bold=True))
    rows.append(["Headcount at year end (incl. 2 founders)"] + [num(yend(b, "headcount", y)) for y in Y] + [""])
    rows.append(["FTE at year end"] + [num(yend(b, "fte", y), 2) for y in Y] + [""])
    t["TABLE_STAFF"] = table(H3T, rows)

    rows = []
    for h in GROWTH.staff:
        rows.append([h.role, label(h.start), f"{h.fte:.2f}", gbp(h.annual_salary)])
    t["TABLE_G_STAFF"] = table(["Role (growth case)", "Paid from", "FTE", "Annual salary"], rows, 2)
    rows = [annual_row(g, "Gross salaries", "salaries"), annual_row(g, "Employer NI (after Employment Allowance)", "employer_ni"),
            annual_row(g, "Employer pension (3%)", "pension"), annual_row(g, "Total staff cost", "staff_total", bold=True)]
    rows.append(["Headcount at year end"] + [num(yend(g, "headcount", y)) for y in Y] + [""])
    t["TABLE_G_STAFF_COST"] = table(H3T, rows)

    # Growth case acquisition summary
    rows = []
    for lbl, key in [("Partner-university licences (year end)", "licences_active"),
                     ("Students covered by licences (year end)", "licence_students"),
                     ("Insights Starter subscriptions (year end)", "starter_subs"),
                     ("Cumulative verified students (year end)", "registered"),
                     ("Premium subscribers (year end)", "premium_subs")]:
        rows.append([lbl] + [num(yend(g, key, y)) for y in Y])
    rows.append(["NHS ICB contracts (year end)", "0", "0", "1"])
    t["TABLE_G_ACQ"] = table(H3, rows)

    # Sensitivity: what if no licence is ever signed?
    downside = Scenario(**{**BASE.__dict__, "name": "Downside", "licences": [], "staff": [], "lean": True,
                           "licensed_share_of_active": (0.0, 0.0, 0.0),
                           "marketing": {k: (v[0], v[0], v[0]) for k, v in BASE_MARKETING.items()},
                           "capex": {}})
    d = run(downside)
    t["DOWN_CASH_Y3"] = gbp(yend(d, "closing_cash", 3))
    t["DOWN_MIN_CASH"] = gbp(min(d["closing_cash"][1:]))
    t["DOWN_REV_Y3"] = gbp(ysum(d, "revenue", 3))

    # Unit economics
    lic_cac = (ysum(b, "cos_onboarding", 2) + ysum(b, "cos_onboarding", 3) - 0) / max(1, len(BASE.licences))
    t["LIC_ONBOARD_COST"] = gbp(600)
    t["PREMIUM_NET"] = f"£{PREMIUM_PRICE * (1 - PREMIUM_FEE_RATE):.2f}"
    t["LIC_ACV"] = gbp(20_000)

    # Annual summary for exec summary
    rows = []
    for lbl, f in [("Total revenue", lambda y: gbp(ysum(b, "revenue", y))),
                   ("Gross margin", lambda y: t[f"GM_Y{y}"]),
                   ("Profit / (loss) after tax", lambda y: t[f"PAT_Y{y}"]),
                   ("Closing cash", lambda y: t[f"CASH_Y{y}"]),
                   ("Paying partner universities (year end)", lambda y: t[f"LIC_END_Y{y}"]),
                   ("Insights Starter subscriptions (year end)", lambda y: t[f"STARTER_END_Y{y}"]),
                   ("Verified students on platform (year end)", lambda y: t[f"REG_Y{y}"]),
                   ("Team headcount / FTE (year end)", lambda y: f"{num(yend(b, 'headcount', y))} / {num(yend(b, 'fte', y), 2)}" + (" (founders unpaid)" if y < 3 else "")),
                   ]:
        rows.append([lbl] + [f(y) for y in Y])
    t["TABLE_HEADLINE"] = table(["Base case", "Year 1 (FY27)", "Year 2 (FY28)", "Year 3 (FY29)"], rows)
    rows = []
    for lbl, f in [("Total revenue", lambda y: gbp(ysum(g, "revenue", y))),
                   ("Operating profit / (loss)", lambda y: gbp(ysum(g, "operating_profit", y))),
                   ("Closing cash", lambda y: gbp(yend(g, "closing_cash", y))),
                   ("Headcount (year end)", lambda y: num(yend(g, "headcount", y)))]:
        rows.append([lbl] + [f(y) for y in Y])
    t["TABLE_G_HEADLINE"] = table(["Growth case", "Year 1 (FY27)", "Year 2 (FY28)", "Year 3 (FY29)"], rows)
    return t


# ---------------------------------------------------------------------------
# Excel workbook
# ---------------------------------------------------------------------------
def write_xlsx(results: dict[str, dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    navy = PatternFill("solid", fgColor="0F2942")
    rows_spec = [
        ("CUSTOMERS", None), ("New verified sign-ups", "new_signups"), ("Cumulative verified students", "registered"),
        ("Monthly active students", "active"), ("Premium subscribers", "premium_subs"),
        ("Insights Starter subscriptions", "starter_subs"), ("Partner licences active", "licences_active"),
        ("Students covered by licences", "licence_students"),
        ("REVENUE (£)", None), ("Partner licences", "rev_licence"), ("Insights Starter", "rev_starter"),
        ("Premium", "rev_premium"), ("NHS ICB", "rev_icb"), ("Total revenue", "revenue"),
        ("COST OF SALES (£)", None), ("Hosting", "cos_hosting"), ("AI inference", "cos_ai"),
        ("Email/push/SMS", "cos_messaging"), ("Payment & app-store fees", "cos_payment"),
        ("Onboarding & pilot", "cos_onboarding"), ("Customer-success tooling", "cos_success"),
        ("Total cost of sales", "cos_total"), ("Gross profit", "gross_profit"),
        ("OVERHEADS (£)", None), ("Salaries", "salaries"), ("Employer NI (net of EA)", "employer_ni"),
        ("Employer pension", "pension"), ("Other staff", "other_staff"), ("Total staff", "staff_total"),
        ("Marketing", "marketing"), ("Software", "software"), ("Insurance", "insurance"),
        ("Accountancy", "accountancy"), ("Legal", "legal"), ("Office", "office"), ("Compliance", "compliance"),
        ("Travel", "travel"), ("Bank & comms", "bank_comms"), ("EBITDA", "ebitda"), ("Depreciation", "depreciation"),
        ("Operating profit", "operating_profit"),
        ("CASH (£)", None), ("Licence/contract invoices raised", "licence_invoiced"), ("Cash receipts", "receipts"),
        ("Capex", "capex"), ("Corporation tax paid", "ct_paid"), ("Total payments", "payments"),
        ("Net cash flow", "net_cash"), ("Closing cash", "closing_cash"),
        ("TEAM", None), ("Headcount", "headcount"), ("FTE", "fte"),
    ]
    for name, r in results.items():
        ws = wb.create_sheet(f"{name} - monthly")
        ws.append(["DEQUAD Ltd — " + name + " (36-month workings, GBP). M1 = Sep 2026; FY ends 31 Aug."])
        ws["A1"].font = Font(bold=True, size=12)
        hdr = ["Line"] + [label(m) for m in range(1, MONTHS + 1)] + ["FY27", "FY28", "FY29"]
        ws.append(hdr)
        for c in range(1, len(hdr) + 1):
            ws.cell(row=2, column=c).font = Font(bold=True, color="FFFFFF")
            ws.cell(row=2, column=c).fill = navy
        for lbl, key in rows_spec:
            if key is None:
                ws.append([lbl])
                ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
                continue
            vals = [round(r[key][m], 2) for m in range(1, MONTHS + 1)]
            stock_like = key in ("registered", "active", "premium_subs", "starter_subs", "licences_active",
                                 "licence_students", "closing_cash", "headcount", "fte")
            yrs = [vals[y * 12 - 1] if stock_like else round(sum(vals[(y - 1) * 12:y * 12]), 2) for y in (1, 2, 3)]
            ws.append([lbl] + vals + yrs)
        ws.column_dimensions["A"].width = 34
        for c in range(2, len(hdr) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 10
            for row in ws.iter_rows(min_row=3, min_col=c, max_col=c):
                for cell in row:
                    cell.number_format = "#,##0;(#,##0);-"
                    cell.alignment = Alignment(horizontal="right")
        ws.freeze_panes = "B3"

        bs = wb.create_sheet(f"{name} - balance sheet")
        bs.append(["At 31 August", "2027", "2028", "2029"])
        for k in ["fixed_cost", "acc_dep", "nbv", "stock", "debtors", "cash", "total_assets", "deferred_income",
                  "ct_creditor", "total_liabilities", "net_assets", "share_capital", "retained", "equity"]:
            bs.append([k.replace("_", " ")] + [round(r["_bs"][y][k], 2) for y in (1, 2, 3)])
        bs.column_dimensions["A"].width = 28

    a = wb.create_sheet("Assumptions", 0)
    for line in [
        ("Premium price (£/month)", PREMIUM_PRICE), ("Insights Starter price (£/month)", STARTER_PRICE),
        ("Partner licence price (£/enrolled student/year)", LICENCE_PRICE_PER_STUDENT),
        ("Opening verified students (beta)", OPENING_REGISTERED), ("Monthly attrition of active students", ACTIVE_MONTHLY_ATTRITION),
        ("Premium billing live from month", PREMIUM_LIVE_FROM), ("Blended Premium fee rate", PREMIUM_FEE_RATE),
        ("Corporation tax (small profits rate)", CORP_TAX), ("Employer NI rate / threshold", f"{NI_RATE} / {NI_THRESHOLD}"),
        ("Employment Allowance", EMPLOYMENT_ALLOWANCE), ("Employer pension", EMPLOYER_PENSION),
        ("NatWest in-kind support until month", NATWEST_IN_KIND_UNTIL), ("Founder share capital (£)", 6000),
        ("Licence billing", "Annual in advance, 30-day terms, revenue recognised monthly"),
        ("VAT", f"Not registered — taxable turnover below £{VAT_THRESHOLD:,} in base case"),
    ]:
        a.append(list(line))
    a.column_dimensions["A"].width = 48
    a.column_dimensions["B"].width = 50
    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX}")


def main() -> None:
    b = run(BASE)
    g = run(GROWTH)
    tokens = build_tokens(b, g)
    text = TEMPLATE_MD.read_text(encoding="utf-8")
    missing = set(re.findall(r"\{\{([A-Z0-9_]+)\}\}", text)) - set(tokens)
    if missing:
        raise SystemExit(f"template tokens with no value: {sorted(missing)}")
    text = re.sub(r"\{\{([A-Z0-9_]+)\}\}", lambda mt: tokens[mt.group(1)], text)
    OUT_MD.write_text(text, encoding="utf-8")
    print(f"wrote {OUT_MD}")
    write_xlsx({"Base case": b, "Growth case": g})

    import build_ukes_pdf as pdf
    chrome = shutil.which("chromium") or "/opt/pw-browsers/chromium"
    pdf.CHROME = chrome
    pdf.CSS += """
    table:has(th:nth-child(10)) { font-size: 6.6pt; table-layout: auto; }
    table:has(th:nth-child(10)) th, table:has(th:nth-child(10)) td { padding: 3px 3px; white-space: nowrap; }
    table:has(th:nth-child(10)) td:first-child, table:has(th:nth-child(10)) th:first-child { white-space: normal; }
    body { max-width: 800px; }
    """
    pdf.build(OUT_MD, HERE / "DEQUAD_UKES_Business_Plan_v6.html", HERE / "DEQUAD_UKES_Business_Plan_v6.pdf",
              "DEQUAD — Business Plan v6.0 (UKES Innovator Founder Endorsement)")


if __name__ == "__main__":
    import sys
    sys.path.insert(0, str(HERE))
    if "--numbers" in sys.argv:
        b, g = run(BASE), run(GROWTH)
        for name, r in (("BASE", b), ("GROWTH", g)):
            print(name)
            for y in (1, 2, 3):
                print(f"  Y{y}: rev {ysum(r,'revenue',y):,.0f} cos {ysum(r,'cos_total',y):,.0f} staff {ysum(r,'staff_total',y):,.0f} "
                      f"ovh {ysum(r,'overheads_ex_staff',y):,.0f} op {ysum(r,'operating_profit',y):,.0f} tax {r['_tax_due'][y]:,.0f} "
                      f"cash {yend(r,'closing_cash',y):,.0f} reg {yend(r,'registered',y):,.0f} prem {yend(r,'premium_subs',y)} "
                      f"start {yend(r,'starter_subs',y)} lic {yend(r,'licences_active',y)}")
            print("  min cash", min(r["closing_cash"][1:]))
    else:
        main()
